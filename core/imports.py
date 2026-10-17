"""Phase 7 bulk import of classes, sections, guardians and students (ADR 0017).

Flow: an administrator uploads one .xlsx workbook for a chosen Academic Year
and enrollment start date. `create_batch` parses and validates it without
changing any school record (a dry run) and stores an ImportBatch with one
ImportRowResult per problem. A batch without errors keeps its parsed rows
and can be committed once, within PENDING_LIFETIME; `commit_batch` validates
the rows again against the current records and then creates everything in
one transaction, through the normal domain workflow: Applicant, accepted
admission decision, `progress_to_student()`, admission number, enrollment and
roll number. Nothing bypasses those rules, and either the whole batch is
imported or nothing is.

Workbook layout (sheet and column names are matched without regard to case,
spaces or underscores):
- "Students" (required): one row per student.
- "Classes" (optional): Class and Section pairs to create if missing. A
  student row may only name a Class/Section that exists already or is listed
  here, so a typo never silently creates a new class.
Any other sheet (for example "Instructions") is ignored.

Messages never contain names, phone numbers or other personal data; they
name the sheet, row and column instead.
"""

import csv
import datetime
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field

from django.core.exceptions import ObjectDoesNotExist, ValidationError
from django.core.validators import EmailValidator
from django.db import IntegrityError, transaction
from django.utils import timezone

from accounts.models import SecurityEvent
from core.models import (
    ADMISSION_NUMBER_VALIDATOR,
    AcademicEnrollment,
    Applicant,
    ApplicantGuardian,
    ClassGrade,
    Guardian,
    ImportBatch,
    ImportRowResult,
    RelationshipType,
    Section,
    Student,
    business_today,
    normalize_admission_number,
)

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_SHEETS = 10
MAX_STUDENT_ROWS = 3000
MAX_CLASS_ROWS = 500
MAX_COLUMNS = 40
PENDING_LIFETIME = datetime.timedelta(hours=24)

SHEET_STUDENTS = "Students"
SHEET_CLASSES = "Classes"
BATCH = "Batch"

CLASS_COLUMNS = [("class", "Class", True), ("section", "Section", True)]
STUDENT_COLUMNS = [
    ("full_name", "Full name", True),
    ("date_of_birth", "Date of birth", False),
    ("admission_number", "Admission number", False),
    ("class", "Class", True),
    ("section", "Section", True),
    ("roll_number", "Roll number", False),
    ("guardian1_name", "Guardian 1 name", False),
    ("guardian1_relationship", "Guardian 1 relationship", False),
    ("guardian1_phone", "Guardian 1 phone", False),
    ("guardian1_email", "Guardian 1 email", False),
    ("guardian2_name", "Guardian 2 name", False),
    ("guardian2_relationship", "Guardian 2 relationship", False),
    ("guardian2_phone", "Guardian 2 phone", False),
    ("guardian2_email", "Guardian 2 email", False),
]
LABELS = {key: label for key, label, _ in CLASS_COLUMNS + STUDENT_COLUMNS}

RELATIONSHIPS = {
    "father": RelationshipType.FATHER,
    "mother": RelationshipType.MOTHER,
    "guardian": RelationshipType.GUARDIAN,
    "grandparent": RelationshipType.GRANDPARENT,
    "grandfather": RelationshipType.GRANDPARENT,
    "grandmother": RelationshipType.GRANDPARENT,
    "other": RelationshipType.OTHER,
}
ROLL_NUMBER_MAX = 32767


class ImportFileError(Exception):
    """The upload cannot be read at all (wrong type, too large, corrupt)."""


@dataclass
class Issue:
    sheet: str
    row: int | None
    column: str
    level: str
    message: str


@dataclass
class Plan:
    """What committing would create; JSON-free, rebuilt on every validation."""

    classes: list = field(default_factory=list)  # new ClassGrade names
    sections: list = field(default_factory=list)  # (class key, section name) new
    guardians: list = field(default_factory=list)  # dicts; existing_id or None
    students: list = field(default_factory=list)  # dicts

    def counts(self):
        return {
            "classes_new": len(self.classes),
            "sections_new": len(self.sections),
            "guardians_new": sum(1 for g in self.guardians if g["existing_id"] is None),
            "guardians_existing": sum(1 for g in self.guardians if g["existing_id"]),
            "students": len(self.students),
        }


# --- Normalisation helpers ---------------------------------------------------


def name_key(value):
    """Comparison key for names: whitespace collapsed, case folded."""
    return " ".join((value or "").split()).casefold()


def header_key(value):
    return re.sub(r"[^0-9a-z]", "", str(value or "").casefold())


def normalize_phone(value):
    """An Indian-style phone number in +<country><number> form, or None.

    Accepts 10-digit mobiles (with optional leading 0 or 91/+91) and other
    numbers written with an explicit leading '+'. Returns "" for blank input
    and None when the value is not a recognisable number.
    """
    raw = (value or "").strip()
    if not raw:
        return ""
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 10 and digits[0] in "6789":
        return "+91" + digits
    if len(digits) == 11 and digits.startswith("0") and digits[1] in "6789":
        return "+91" + digits[1:]
    if len(digits) == 12 and digits.startswith("91") and digits[2] in "6789":
        return "+" + digits
    if raw.startswith("+") and 8 <= len(digits) <= 15:
        return "+" + digits
    return None


def parse_date(value):
    """ISO (YYYY-MM-DD) or Indian day-first (DD/MM/YYYY, DD-MM-YYYY, DD.MM.YYYY)."""
    value = (value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValueError(value)


def _cell_text(value):
    """Cell value as clean text; dates become ISO strings."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int):
        return str(value)
    return " ".join(str(value).split())


def neutralize_csv(value):
    """Stop spreadsheet programs from treating exported text as a formula."""
    text = str(value)
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


# --- Reading the workbook ----------------------------------------------------


def read_workbook(data):
    """Return ({sheet: [(row_number, {key: text})]}, [Issue]) from .xlsx bytes."""
    import openpyxl

    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError("The file is larger than 2 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise ImportFileError("Workbooks with macros are not accepted.")
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ImportFileError("The workbook is too large to import.")
    except zipfile.BadZipFile:
        raise ImportFileError("This is not an Excel .xlsx workbook.") from None
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception:  # openpyxl raises many types for damaged files
        raise ImportFileError("The workbook could not be read.") from None

    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise ImportFileError("The workbook has too many sheets.")
        by_key = {header_key(name): name for name in workbook.sheetnames}
        if header_key(SHEET_STUDENTS) not in by_key:
            raise ImportFileError('The workbook has no "Students" sheet.')
        issues = []
        sheets = {}
        for sheet, columns, limit in (
            (SHEET_STUDENTS, STUDENT_COLUMNS, MAX_STUDENT_ROWS),
            (SHEET_CLASSES, CLASS_COLUMNS, MAX_CLASS_ROWS),
        ):
            actual = by_key.get(header_key(sheet))
            if actual is None:
                sheets[sheet] = []
                continue
            sheets[sheet] = _read_sheet(workbook[actual], sheet, columns, limit, issues)
        return sheets, issues
    finally:
        workbook.close()


def _read_sheet(worksheet, sheet, columns, limit, issues):
    rows = worksheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)
    header = next(rows, None)
    if header is None:
        issues.append(Issue(sheet, None, "", "ERROR", "The sheet is empty."))
        return []
    known = {header_key(label): key for key, label, _ in columns}
    known.update({header_key(key): key for key, _, _ in columns})
    positions = {}
    for index, title in enumerate(header):
        key = known.get(header_key(title))
        if key is None:
            if _cell_text(title):
                issues.append(Issue(sheet, 1, _cell_text(title)[:50], "WARNING",
                                    "Unknown column ignored."))
            continue
        if key in positions:
            issues.append(Issue(sheet, 1, LABELS[key], "ERROR", "This column appears twice."))
            continue
        positions[key] = index
    for key, label, required in columns:
        if required and key not in positions:
            issues.append(Issue(sheet, 1, label, "ERROR", "Required column is missing."))

    parsed = []
    for row_number, values in enumerate(rows, start=2):
        if values is None or all(_cell_text(v) == "" for v in values):
            continue
        if len(parsed) >= limit:
            issues.append(Issue(sheet, row_number, "", "ERROR",
                                f"Too many rows; at most {limit} are allowed per import."))
            break
        record = {}
        for key, index in positions.items():
            value = values[index] if index < len(values) else None
            if isinstance(value, str) and value.lstrip().startswith("="):
                issues.append(Issue(sheet, row_number, LABELS[key], "ERROR",
                                    "Formulas are not allowed; paste values only."))
                value = None
            record[key] = _cell_text(value)
        parsed.append([row_number, record])
    return parsed


# --- Validation (no writes) --------------------------------------------------


class _Context:
    """Current school records needed to validate, loaded once."""

    def __init__(self, academic_year):
        self.classes = {name_key(c.name): c for c in ClassGrade.objects.all()}
        self.sections = {
            (s.class_grade_id, name_key(s.name)): s for s in Section.objects.all()
        }
        self.guardians_by_identity = {}
        self.guardian_phones = set()
        for guardian in Guardian.objects.order_by("pk"):
            phone = normalize_phone(guardian.phone)
            if phone:
                self.guardian_phones.add(phone)
                self.guardians_by_identity.setdefault(
                    (phone, name_key(guardian.full_name)), guardian
                )
        self.admission_numbers = {
            value.upper()
            for value in Student.objects.exclude(admission_number__isnull=True)
            .values_list("admission_number", flat=True)
        }
        self.student_identities = {
            (name_key(name), dob)
            for name, dob in Student.objects.exclude(date_of_birth__isnull=True)
            .values_list("full_name", "date_of_birth")
        }
        self.roll_numbers = set(
            AcademicEnrollment.objects.filter(
                academic_year=academic_year, roll_number__isnull=False
            ).values_list("section_id", "roll_number")
        )


def validate(sheets, academic_year, enrollment_start_date):
    """Return (Plan, [Issue]) for parsed rows. Never writes."""
    issues = []
    plan = Plan()

    def error(sheet, row, column, message):
        issues.append(Issue(sheet, row, column, "ERROR", message))

    def warn(sheet, row, column, message):
        issues.append(Issue(sheet, row, column, "WARNING", message))

    start = enrollment_start_date
    if isinstance(start, str):
        start = datetime.date.fromisoformat(start)
    if not academic_year.contains(start):
        error(BATCH, None, "", "The enrollment start date is outside the Academic Year.")
    if start > business_today():
        error(BATCH, None, "", "The enrollment start date cannot be in the future.")

    context = _Context(academic_year)

    # Classes sheet: explicit list of Class/Section pairs that may be created.
    new_classes = {}  # class key -> display name
    allowed_sections = {}  # (class key, section key) -> (class name, section name)
    for row, record in sheets.get(SHEET_CLASSES, []):
        class_name = record.get("class", "")
        section_name = record.get("section", "")
        if not class_name:
            error(SHEET_CLASSES, row, "Class", "Class is required.")
        if not section_name:
            error(SHEET_CLASSES, row, "Section", "Section is required.")
        if not class_name or not section_name:
            continue
        if len(class_name) > 100 or len(section_name) > 100:
            error(SHEET_CLASSES, row, "", "Class and Section names are limited to 100 characters.")
            continue
        class_key, section_key = name_key(class_name), name_key(section_name)
        if (class_key, section_key) in allowed_sections:
            warn(SHEET_CLASSES, row, "", "This Class and Section is listed twice.")
            continue
        existing_class = context.classes.get(class_key)
        if existing_class is None:
            new_classes.setdefault(class_key, class_name)
            plan.sections.append((class_key, section_name))
        elif (existing_class.pk, section_key) not in context.sections:
            plan.sections.append((class_key, section_name))
        allowed_sections[(class_key, section_key)] = (class_name, section_name)
    plan.classes = list(new_classes.values())

    def resolve_section(class_name, section_name):
        """'existing' Section, ('new', class key, section key) or None."""
        class_key, section_key = name_key(class_name), name_key(section_name)
        existing_class = context.classes.get(class_key)
        if existing_class is not None:
            existing = context.sections.get((existing_class.pk, section_key))
            if existing is not None:
                return existing
        if (class_key, section_key) in allowed_sections:
            return ("new", class_key, section_key)
        return None

    guardians = {}  # identity -> index into plan.guardians
    file_admission_numbers = {}
    file_roll_numbers = {}
    file_identities = {}
    file_names = {}

    for row, record in sheets.get(SHEET_STUDENTS, []):
        ok = True
        full_name = record.get("full_name", "")
        if not full_name:
            error(SHEET_STUDENTS, row, "Full name", "Full name is required.")
            ok = False
        elif len(full_name) > 255:
            error(SHEET_STUDENTS, row, "Full name", "Full name is longer than 255 characters.")
            ok = False

        date_of_birth = None
        if record.get("date_of_birth"):
            try:
                date_of_birth = parse_date(record["date_of_birth"])
            except ValueError:
                error(SHEET_STUDENTS, row, "Date of birth",
                      "Use a date cell, YYYY-MM-DD or DD/MM/YYYY.")
                ok = False
            else:
                if date_of_birth > business_today():
                    error(SHEET_STUDENTS, row, "Date of birth",
                          "Date of birth cannot be in the future.")
                    ok = False

        admission_number = normalize_admission_number(record.get("admission_number"))
        if admission_number is not None:
            try:
                ADMISSION_NUMBER_VALIDATOR(admission_number)
            except ValidationError:
                error(SHEET_STUDENTS, row, "Admission number",
                      "Use letters, digits, slash, dot, underscore or hyphen (up to 30).")
                ok = False
            else:
                key = admission_number.upper()
                if key in context.admission_numbers:
                    error(SHEET_STUDENTS, row, "Admission number",
                          "Another Student already has this admission number.")
                    ok = False
                elif key in file_admission_numbers:
                    error(SHEET_STUDENTS, row, "Admission number",
                          f"Same admission number as row {file_admission_numbers[key]}.")
                    ok = False
                else:
                    file_admission_numbers[key] = row

        section = None
        class_name, section_name = record.get("class", ""), record.get("section", "")
        if not class_name or not section_name:
            error(SHEET_STUDENTS, row, "Class" if not class_name else "Section",
                  "Class and Section are required.")
            ok = False
        else:
            section = resolve_section(class_name, section_name)
            if section is None:
                error(SHEET_STUDENTS, row, "Section",
                      f'Class "{class_name[:100]}" Section "{section_name[:100]}" does not '
                      "exist. Create it first or list it on the Classes sheet.")
                ok = False

        roll_number = None
        if record.get("roll_number"):
            text = record["roll_number"]
            if not text.isdigit() or not 1 <= int(text) <= ROLL_NUMBER_MAX:
                error(SHEET_STUDENTS, row, "Roll number", "Roll number must be a positive "
                      "whole number.")
                ok = False
            elif section is not None:
                roll_number = int(text)
                section_ref = section.pk if isinstance(section, Section) else section
                taken = isinstance(section, Section) and (
                    (section.pk, roll_number) in context.roll_numbers
                )
                if taken:
                    error(SHEET_STUDENTS, row, "Roll number", "This roll number is already used "
                          "in this Section for this Academic Year.")
                    ok = False
                elif (section_ref, roll_number) in file_roll_numbers:
                    error(SHEET_STUDENTS, row, "Roll number", "Same roll number and Section as "
                          f"row {file_roll_numbers[(section_ref, roll_number)]}.")
                    ok = False
                else:
                    file_roll_numbers[(section_ref, roll_number)] = row

        links = []
        seen_here = set()
        for number in (1, 2):
            prefix = f"guardian{number}_"
            values = {k: record.get(prefix + k, "") for k in ("name", "relationship", "phone",
                                                                "email")}
            label = f"Guardian {number}"
            if not any(values.values()):
                # A guardian is optional (Phase 1 lets a Student exist with none);
                # once any of its cells is filled, name and relationship are required.
                continue
            guardian_ok = True
            if not values["name"]:
                error(SHEET_STUDENTS, row, f"{label} name", f"{label} name is required.")
                guardian_ok = False
            elif len(values["name"]) > 255:
                error(SHEET_STUDENTS, row, f"{label} name", "Name is longer than 255 characters.")
                guardian_ok = False
            relationship = RELATIONSHIPS.get(values["relationship"].casefold())
            if relationship is None:
                error(SHEET_STUDENTS, row, f"{label} relationship",
                      "Use Father, Mother, Guardian, Grandparent or Other.")
                guardian_ok = False
            phone = normalize_phone(values["phone"])
            if phone is None:
                error(SHEET_STUDENTS, row, f"{label} phone",
                      "Phone number not recognised; use a 10-digit mobile or +country code.")
                guardian_ok = False
            email = values["email"]
            if email:
                try:
                    EmailValidator()(email)
                except ValidationError:
                    error(SHEET_STUDENTS, row, f"{label} email", "Email address is not valid.")
                    guardian_ok = False
            if not guardian_ok:
                ok = False
                continue

            key_name = name_key(values["name"])
            if phone:
                identity = ("phone", phone, key_name)
            elif email:
                identity = ("email", email.casefold(), key_name)
            else:
                identity = ("row", row, number)
                warn(SHEET_STUDENTS, row, f"{label} phone",
                     "No phone or email: this guardian cannot be matched to siblings or "
                     "receive messages.")
            if identity in seen_here:
                error(SHEET_STUDENTS, row, f"{label} name",
                      "The same guardian is listed twice for this student.")
                ok = False
                continue
            seen_here.add(identity)
            if identity not in guardians:
                existing = None
                if phone:
                    existing = context.guardians_by_identity.get((phone, key_name))
                    if existing is None and phone in context.guardian_phones:
                        warn(SHEET_STUDENTS, row, f"{label} phone",
                             "This phone belongs to an existing guardian with a different "
                             "name; a new guardian record will be created. Check before "
                             "importing.")
                guardians[identity] = len(plan.guardians)
                plan.guardians.append({
                    "full_name": values["name"],
                    "phone": phone,
                    "email": email,
                    "existing_id": existing.pk if existing else None,
                    "first_row": row,
                })
            links.append({"guardian": guardians[identity], "relationship": relationship.value})

        if ok and not links:
            warn(SHEET_STUDENTS, row, "Guardian 1 name",
                 "No guardian listed: this student will have no guardian, so no portal "
                 "access or messages until one is added.")

        if full_name:
            identity = (name_key(full_name), date_of_birth)
            if date_of_birth and identity in context.student_identities:
                warn(SHEET_STUDENTS, row, "Full name", "A Student with the same name and date "
                     "of birth already exists. Check this is not a duplicate.")
            if date_of_birth and identity in file_identities:
                warn(SHEET_STUDENTS, row, "Full name", "Same name and date of birth as row "
                     f"{file_identities[identity]}.")
            elif name_key(full_name) in file_names:
                warn(SHEET_STUDENTS, row, "Full name",
                     f"Same name as row {file_names[name_key(full_name)]}.")
            file_identities.setdefault(identity, row)
            file_names.setdefault(name_key(full_name), row)

        if ok:
            if isinstance(section, Section):
                placement = str(section)
            else:
                placement = "{} - {} (new)".format(*allowed_sections[(section[1], section[2])])
            plan.students.append({
                "row": row,
                "placement": placement,
                "full_name": full_name,
                "date_of_birth": date_of_birth,
                "admission_number": admission_number,
                "section": section,
                "roll_number": roll_number,
                "links": links,
            })

    if not sheets.get(SHEET_STUDENTS):
        error(SHEET_STUDENTS, None, "", "The Students sheet has no rows to import.")
    return plan, issues


# --- Batches -----------------------------------------------------------------


def expire_stale_batches(now=None):
    """Move VALIDATED batches older than PENDING_LIFETIME to EXPIRED."""
    now = now or timezone.now()
    stale = ImportBatch.objects.filter(
        status=ImportBatch.Status.VALIDATED, uploaded_at__lt=now - PENDING_LIFETIME
    )
    return stale.update(status=ImportBatch.Status.EXPIRED, payload={}, finished_at=now)


def create_batch(*, user, file_name, data, academic_year, enrollment_start_date):
    """Validate an upload (dry run) and record it as an ImportBatch.

    Raises ImportFileError when the file cannot be read at all.
    """
    sheets, issues = read_workbook(data)
    plan, more = validate(sheets, academic_year, enrollment_start_date)
    issues += more
    errors = sum(1 for issue in issues if issue.level == "ERROR")
    warnings = len(issues) - errors
    status = ImportBatch.Status.REJECTED if errors else ImportBatch.Status.VALIDATED
    with transaction.atomic():
        batch = ImportBatch.objects.create(
            uploaded_by=user,
            file_name=(file_name or "upload.xlsx")[:255],
            file_sha256=hashlib.sha256(data).hexdigest(),
            academic_year=academic_year,
            enrollment_start_date=enrollment_start_date,
            student_rows=len(sheets.get(SHEET_STUDENTS, [])),
            class_rows=len(sheets.get(SHEET_CLASSES, [])),
            error_count=errors,
            warning_count=warnings,
            status=status,
            payload=sheets if not errors else {},
            summary=plan.counts() if not errors else {},
            finished_at=timezone.now() if errors else None,
        )
        ImportRowResult.objects.bulk_create(
            ImportRowResult(
                batch=batch,
                sheet=issue.sheet,
                row_number=issue.row,
                column=issue.column[:50],
                level=issue.level,
                message=issue.message[:500],
            )
            for issue in issues
        )
    return batch


def preview(batch):
    """Re-validate a pending batch against current records: (Plan, [Issue])."""
    return validate(batch.payload, batch.academic_year, batch.enrollment_start_date)


class CommitError(Exception):
    pass


REVIEW_WARNINGS = "Confirm that you have reviewed the warnings before importing."


def commit_batch(batch, *, user, warnings_reviewed=False):
    """Import a VALIDATED batch atomically. Raises CommitError on failure.

    Warnings are checked again against current records, so a warning that
    appeared after the upload still needs `warnings_reviewed`. A missing
    review leaves the batch VALIDATED.
    """
    expire_stale_batches()
    try:
        with transaction.atomic():
            locked = ImportBatch.objects.select_for_update().get(pk=batch.pk)
            if locked.status != ImportBatch.Status.VALIDATED:
                raise CommitError("This import is no longer waiting for confirmation.")
            plan, issues = preview(locked)
            if any(issue.level == "ERROR" for issue in issues):
                raise _Failed("School records changed since the file was checked; "
                              "upload the file again to see the problems.")
            if any(issue.level == "WARNING" for issue in issues) and not warnings_reviewed:
                raise _NeedsReview(REVIEW_WARNINGS)
            counts = _apply(locked, plan, user)
            ImportBatch.objects.filter(pk=locked.pk).update(
                status=ImportBatch.Status.COMMITTED,
                payload={},
                summary=counts,
                finished_at=timezone.now(),
                finished_by=user,
            )
            SecurityEvent.record(
                SecurityEvent.EventType.IMPORT_COMMITTED,
                actor=user,
                detail={"batch": locked.pk, **counts},
            )
    except _NeedsReview as problem:
        raise CommitError(str(problem)) from None
    except _Failed as failure:
        _mark_failed(batch, user, str(failure))
        raise CommitError(str(failure)) from None
    except (ValidationError, IntegrityError, ObjectDoesNotExist):
        reason = "The import could not be completed and nothing was imported."
        _mark_failed(batch, user, reason)
        raise CommitError(reason) from None
    batch.refresh_from_db()
    return batch


class _Failed(Exception):
    pass


class _NeedsReview(Exception):
    pass


def _mark_failed(batch, user, reason):
    ImportBatch.objects.filter(pk=batch.pk, status=ImportBatch.Status.VALIDATED).update(
        status=ImportBatch.Status.FAILED,
        payload={},
        failure_reason=reason[:500],
        finished_at=timezone.now(),
        finished_by=user,
    )


def discard_batch(batch, *, user):
    return ImportBatch.objects.filter(pk=batch.pk, status=ImportBatch.Status.VALIDATED).update(
        status=ImportBatch.Status.DISCARDED,
        payload={},
        finished_at=timezone.now(),
        finished_by=user,
    )


def _apply(batch, plan, user):
    """Create everything in `plan` through the domain workflow."""
    classes = {name_key(c.name): c for c in ClassGrade.objects.all()}
    for name in plan.classes:
        if name_key(name) not in classes:
            classes[name_key(name)] = ClassGrade.objects.create(name=name)
    new_sections = {}
    for class_key, section_name in plan.sections:
        class_grade = classes[class_key]
        section = Section.objects.filter(
            class_grade=class_grade, name__iexact=section_name
        ).first() or Section.objects.create(class_grade=class_grade, name=section_name)
        new_sections[(class_key, name_key(section_name))] = section

    guardians = []
    created_guardians = 0
    for entry in plan.guardians:
        if entry["existing_id"]:
            guardians.append(Guardian.objects.get(pk=entry["existing_id"]))
        else:
            guardians.append(Guardian.objects.create(
                full_name=entry["full_name"], phone=entry["phone"], email=entry["email"]
            ))
            created_guardians += 1

    academic_year = batch.academic_year
    for entry in plan.students:
        section = entry["section"]
        if not isinstance(section, Section):
            section = new_sections[(section[1], section[2])]
        applicant = Applicant.objects.create(
            full_name=entry["full_name"],
            date_of_birth=entry["date_of_birth"],
            import_batch=batch,
        )
        ApplicantGuardian.objects.bulk_create(
            ApplicantGuardian(
                applicant=applicant,
                guardian=guardians[link["guardian"]],
                relationship_type=link["relationship"],
            )
            for link in entry["links"]
        )
        applicant.record_admission_decision(Applicant.AdmissionStatus.ACCEPTED)
        student = applicant.progress_to_student()
        if entry["admission_number"]:
            student.set_admission_number(entry["admission_number"])
        enrollment = student.enroll(
            academic_year=academic_year,
            class_grade=section.class_grade,
            section=section,
            start_date=batch.enrollment_start_date,
            actor=user,
        )
        if entry["roll_number"]:
            enrollment.set_roll_number(entry["roll_number"])

    return {
        "classes_new": len(plan.classes),
        "sections_new": len(new_sections),
        "guardians_new": created_guardians,
        "guardians_existing": len(guardians) - created_guardians,
        "students": len(plan.students),
    }


# --- Files for download ------------------------------------------------------


def template_workbook():
    """An empty import workbook with instructions, as .xlsx bytes."""
    import openpyxl
    from openpyxl.styles import Font

    workbook = openpyxl.Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    lines = [
        "School ERP bulk import",
        "",
        "Students sheet: one row per student. Required: Full name, Class, Section.",
        "Guardians are optional (up to two per student). If you fill in any cell for a "
        "guardian, that guardian's name and relationship are required.",
        "Classes sheet (optional): Class and Section pairs to create if they do not exist yet.",
        "A student row may only use a Class/Section that already exists or is listed on "
        "the Classes sheet.",
        "Dates: use a date cell, or YYYY-MM-DD, or DD/MM/YYYY.",
        "Relationship: Father, Mother, Guardian, Grandparent or Other.",
        "Phone: 10-digit mobile number; it is saved as +91XXXXXXXXXX.",
        "Guardians with the same phone and name are shared between siblings and matched to "
        "existing guardians. A guardian with only an email is always added as a new record.",
        "Do not use formulas. Do not add example rows to the Students sheet.",
        "Choose the Academic Year and the enrollment start date when you upload.",
        "Nothing is saved until you review the preview and confirm.",
    ]
    for index, line in enumerate(lines, start=1):
        instructions.cell(row=index, column=1, value=line)
    instructions["A1"].font = Font(bold=True, size=14)
    instructions.column_dimensions["A"].width = 110

    for title, columns in ((SHEET_STUDENTS, STUDENT_COLUMNS), (SHEET_CLASSES, CLASS_COLUMNS)):
        sheet = workbook.create_sheet(title)
        for index, (_, label, required) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index, value=label)
            cell.font = Font(bold=True, color="000000" if required else "555555")
            sheet.column_dimensions[cell.column_letter].width = max(14, len(label) + 4)
        sheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def issues_csv(batch):
    """The batch's row results as CSV text."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Sheet", "Row", "Column", "Level", "Message"])
    for result in batch.row_results.all():
        writer.writerow([
            neutralize_csv(result.sheet),
            result.row_number or "",
            neutralize_csv(result.column),
            result.level,
            neutralize_csv(result.message),
        ])
    return output.getvalue()
