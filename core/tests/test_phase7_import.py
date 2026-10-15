"""Phase 7 Gates 8-9: bulk import through applicant progression (ADR 0017)."""

import datetime
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, transaction
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import SecurityEvent
from core import imports
from core.models import (
    AcademicEnrollment,
    AcademicYear,
    Applicant,
    ApplicantGuardian,
    ClassGrade,
    Guardian,
    ImportBatch,
    ImportRowResult,
    RelationshipType,
    Section,
    Student,
    StudentGuardian,
    business_today,
)

User = get_user_model()

STUDENT_HEADER = [label for _, label, _ in imports.STUDENT_COLUMNS]


def student_row(name, *, dob="", adm="", cls="Class 5", sec="A", roll="",
                g1=("Ravi Kumar", "Father", "9876543210", ""), g2=None):
    row = [name, dob, adm, cls, sec, roll, *g1]
    row += list(g2) if g2 else ["", "", "", ""]
    return row


def workbook(students, classes=None, *, student_header=None):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Students"
    sheet.append(student_header or STUDENT_HEADER)
    for row in students:
        sheet.append(row)
    if classes is not None:
        class_sheet = book.create_sheet("Classes")
        class_sheet.append(["Class", "Section"])
        for row in classes:
            class_sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class ImportFixture(TestCase):
    def setUp(self):
        super().setUp()
        today = business_today()
        self.year = AcademicYear.objects.create(
            name="2026-27",
            start_date=today - datetime.timedelta(days=60),
            end_date=today + datetime.timedelta(days=300),
        )
        self.start = today - datetime.timedelta(days=30)
        self.class5 = ClassGrade.objects.create(name="Class 5")
        self.section_a = Section.objects.create(class_grade=self.class5, name="A")
        self.admin = User.objects.create_user(username="admin1", password="pass-12345")
        self.admin.groups.add(Group.objects.get(name="School Administrators"))
        self.client.login(username="admin1", password="pass-12345")

    def batch_for(self, data, start=None):
        return imports.create_batch(
            user=self.admin, file_name="students.xlsx", data=data,
            academic_year=self.year, enrollment_start_date=start or self.start,
        )

    def upload(self, data, start=None, name="students.xlsx"):
        return self.client.post(reverse("core:import-upload"), {
            "academic_year": self.year.pk,
            "enrollment_start_date": (start or self.start).isoformat(),
            "file": SimpleUploadedFile(name, data),
        })

    def messages_of(self, batch, level="ERROR"):
        return list(batch.row_results.filter(level=level).values_list("message", flat=True))


class HelperTests(TestCase):
    def test_phone_normalisation(self):
        self.assertEqual(imports.normalize_phone("98765 43210"), "+919876543210")
        self.assertEqual(imports.normalize_phone("098765-43210"), "+919876543210")
        self.assertEqual(imports.normalize_phone("+91 98765 43210"), "+919876543210")
        self.assertEqual(imports.normalize_phone("919876543210"), "+919876543210")
        self.assertEqual(imports.normalize_phone("+44 20 7946 0958"), "+442079460958")
        self.assertEqual(imports.normalize_phone(""), "")
        self.assertIsNone(imports.normalize_phone("12345"))
        self.assertIsNone(imports.normalize_phone("not a phone"))

    def test_dates_are_day_first(self):
        self.assertEqual(imports.parse_date("05/06/2015"), datetime.date(2015, 6, 5))
        self.assertEqual(imports.parse_date("2015-06-05"), datetime.date(2015, 6, 5))
        self.assertEqual(imports.parse_date("05.06.2015"), datetime.date(2015, 6, 5))
        with self.assertRaises(ValueError):
            imports.parse_date("June 5")

    def test_csv_cells_are_neutralised(self):
        for value in ("=1+1", "+cmd", "-2", "@SUM(A1)"):
            self.assertTrue(imports.neutralize_csv(value).startswith("'"))
        self.assertEqual(imports.neutralize_csv("Class 5"), "Class 5")


class DryRunTests(ImportFixture):
    def test_validation_changes_no_school_records(self):
        data = workbook([student_row("Asha", roll="1")], classes=[["Class 6", "B"]])
        batch = self.batch_for(data)
        self.assertEqual(batch.status, ImportBatch.Status.VALIDATED)
        self.assertEqual(Student.objects.count(), 0)
        self.assertEqual(Applicant.objects.count(), 0)
        self.assertEqual(Guardian.objects.count(), 0)
        self.assertFalse(ClassGrade.objects.filter(name="Class 6").exists())
        self.assertEqual(batch.summary["students"], 1)
        self.assertEqual(batch.summary["classes_new"], 1)
        self.assertTrue(batch.payload)

    def test_unlisted_class_is_an_error_not_a_new_class(self):
        batch = self.batch_for(workbook([student_row("Asha", cls="Clas 5")]))
        self.assertEqual(batch.status, ImportBatch.Status.REJECTED)
        self.assertEqual(batch.payload, {})
        self.assertTrue(any("does not exist" in m for m in self.messages_of(batch)))

    def test_class_and_section_names_match_without_case(self):
        batch = self.batch_for(workbook([student_row("Asha", cls="class 5", sec="a")]))
        self.assertEqual(batch.status, ImportBatch.Status.VALIDATED)
        self.assertEqual(batch.summary["classes_new"], 0)

    def test_row_errors_are_reported_without_personal_data(self):
        existing = Applicant.objects.create(full_name="Old", admission_status="ACCEPTED")
        student = existing.progress_to_student()
        student.set_admission_number("ADM-1")
        future = (business_today() + datetime.timedelta(days=5)).isoformat()
        rows = [
            student_row("Bhanu Prakash", adm="adm-1"),
            student_row("Chitra Devi", roll="3"),
            student_row("Deepa Rani", roll="3"),
            student_row("Eswar Rao", g1=("Lakshmi", "Aunt", "9876543210", "")),
            student_row("Farhan Ali", g1=("Imran", "Father", "12345", "")),
            student_row("Gita Sree", dob=future),
            student_row("=HYPERLINK(\"x\")"),
            student_row("", g1=("", "", "", "")),
        ]
        batch = self.batch_for(workbook(rows))
        self.assertEqual(batch.status, ImportBatch.Status.REJECTED)
        found = " | ".join(self.messages_of(batch))
        for expected in ("admission number", "Same roll number", "Use Father", "Phone number",
                         "future", "Formulas", "Full name is required"):
            self.assertIn(expected, found)
        for private in ("Bhanu", "Chitra", "Lakshmi", "9876543210", "Imran"):
            self.assertNotIn(private, found)

    def test_missing_required_column(self):
        header = [h for h in STUDENT_HEADER if h != "Guardian 1 relationship"]
        batch = self.batch_for(workbook([], student_header=header))
        self.assertIn("Required column is missing.", self.messages_of(batch))

    def test_empty_template_is_rejected(self):
        batch = self.batch_for(imports.template_workbook())
        self.assertEqual(batch.status, ImportBatch.Status.REJECTED)
        self.assertIn("The Students sheet has no rows to import.", self.messages_of(batch))

    def test_unreadable_files(self):
        with self.assertRaises(imports.ImportFileError):
            self.batch_for(b"not a zip file")
        response = self.upload(b"name,class\n", name="students.csv")
        self.assertContains(response, "Upload an Excel .xlsx workbook.")
        response = self.upload(b"garbage", name="students.xlsx")
        self.assertContains(response, "This is not an Excel .xlsx workbook.")
        self.assertEqual(ImportBatch.objects.count(), 0)

    def test_start_date_must_be_in_year_and_not_future(self):
        future = business_today() + datetime.timedelta(days=1)
        response = self.upload(workbook([student_row("Asha")]), start=future)
        self.assertContains(response, "cannot be in the future")
        before = self.year.start_date - datetime.timedelta(days=1)
        response = self.upload(workbook([student_row("Asha")]), start=before)
        self.assertContains(response, "must fall within the Academic Year")


class CommitTests(ImportFixture):
    def test_full_import_goes_through_applicant_progression(self):
        existing = Guardian.objects.create(full_name="Sita Devi", phone="98480 22338")
        rows = [
            student_row("Asha", dob="05/06/2015", adm="A-101", roll="1",
                        g2=("Sita Devi", "Mother", "+91 98480 22338", "")),
            student_row("Arjun", dob="2017-01-02", adm="A-102", roll="2",
                        g2=("sita devi", "Mother", "9848022338", "")),
            student_row("Kiran", cls="Class 6", sec="B", roll="1",
                        g1=("Meena", "Guardian", "", "meena@example.com")),
        ]
        batch = self.batch_for(workbook(rows, classes=[["Class 6", "B"]]))
        self.assertEqual(batch.status, ImportBatch.Status.VALIDATED, self.messages_of(batch))

        response = self.client.post(reverse("core:import-commit", args=[batch.pk]))
        self.assertRedirects(response, reverse("core:import-detail", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.COMMITTED)
        self.assertEqual(batch.payload, {})
        self.assertEqual(batch.summary["students"], 3)
        self.assertEqual(batch.finished_by, self.admin)

        self.assertEqual(Student.objects.count(), 3)
        for applicant in Applicant.objects.all():
            self.assertEqual(applicant.import_batch, batch)
            self.assertEqual(applicant.admission_status, Applicant.AdmissionStatus.ACCEPTED)
            self.assertTrue(hasattr(applicant, "student"))

        asha = Student.objects.get(full_name="Asha")
        self.assertEqual(asha.date_of_birth, datetime.date(2015, 6, 5))
        self.assertEqual(asha.admission_number, "A-101")
        enrollment = asha.current_enrollment
        self.assertEqual(enrollment.section, self.section_a)
        self.assertEqual(enrollment.start_date, self.start)
        self.assertEqual(enrollment.roll_number, 1)
        self.assertEqual(enrollment.academic_year, self.year)

        # Siblings share one father record; the mother matches the existing guardian.
        father = Guardian.objects.get(full_name="Ravi Kumar")
        self.assertEqual(father.phone, "+919876543210")
        self.assertEqual(father.student_links.count(), 2)
        self.assertEqual(Guardian.objects.filter(full_name__iexact="Sita Devi").count(), 1)
        self.assertEqual(existing.student_links.count(), 2)
        self.assertTrue(StudentGuardian.objects.filter(
            guardian=existing, relationship_type=RelationshipType.MOTHER).exists())
        self.assertEqual(ApplicantGuardian.objects.count(), 5)

        kiran = Student.objects.get(full_name="Kiran")
        self.assertEqual(kiran.current_enrollment.section.class_grade.name, "Class 6")
        self.assertEqual(kiran.current_enrollment.section.name, "B")

        event = SecurityEvent.objects.get(event_type=SecurityEvent.EventType.IMPORT_COMMITTED)
        self.assertEqual(event.actor, self.admin)
        self.assertEqual(event.detail["batch"], batch.pk)
        self.assertEqual(event.event_class, SecurityEvent.EventClass.HISTORY)

    def test_records_changed_after_check_fail_the_whole_batch(self):
        batch = self.batch_for(workbook([
            student_row("Asha", adm="A-1"), student_row("Arjun", adm="A-2"),
        ]))
        other = Applicant.objects.create(full_name="Other", admission_status="ACCEPTED")
        other.progress_to_student().set_admission_number("A-2")

        with self.assertRaises(imports.CommitError):
            imports.commit_batch(batch, user=self.admin)
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.FAILED)
        self.assertEqual(batch.payload, {})
        self.assertEqual(Student.objects.count(), 1)
        self.assertFalse(Applicant.objects.filter(import_batch=batch).exists())

    def test_domain_failure_rolls_back_everything(self):
        batch = self.batch_for(workbook([student_row("Asha"), student_row("Arjun")]))
        original = Student.enroll
        calls = []

        def failing_enroll(student, **kwargs):
            calls.append(student)
            if len(calls) == 2:
                from django.core.exceptions import ValidationError
                raise ValidationError("boom")
            return original(student, **kwargs)

        Student.enroll = failing_enroll
        try:
            with self.assertRaises(imports.CommitError):
                imports.commit_batch(batch, user=self.admin)
        finally:
            Student.enroll = original
        self.assertEqual(Student.objects.count(), 0)
        self.assertEqual(Guardian.objects.count(), 0)
        self.assertEqual(AcademicEnrollment.objects.count(), 0)
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.FAILED)

    def test_a_batch_commits_only_once(self):
        batch = self.batch_for(workbook([student_row("Asha")]))
        imports.commit_batch(batch, user=self.admin)
        with self.assertRaises(imports.CommitError):
            imports.commit_batch(batch, user=self.admin)
        self.assertEqual(Student.objects.count(), 1)

    def test_warnings_need_explicit_review(self):
        Guardian.objects.create(full_name="Someone Else", phone="9876543210")
        batch = self.batch_for(workbook([student_row("Asha")]))
        self.assertEqual(batch.status, ImportBatch.Status.VALIDATED)
        self.assertEqual(batch.warning_count, 1)
        self.client.post(reverse("core:import-commit", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.VALIDATED)
        self.client.post(reverse("core:import-commit", args=[batch.pk]),
                         {"reviewed_warnings": "yes"})
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.COMMITTED)
        self.assertEqual(Guardian.objects.filter(phone="+919876543210").count(), 1)

    def test_warnings_that_appear_after_the_check_need_review(self):
        batch = self.batch_for(workbook([student_row("Asha")]))
        self.assertEqual(batch.warning_count, 0)
        Guardian.objects.create(full_name="Someone Else", phone="9876543210")

        page = self.client.get(reverse("core:import-detail", args=[batch.pk]))
        self.assertContains(page, 'name="reviewed_warnings"')
        self.client.post(reverse("core:import-commit", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.VALIDATED)
        self.assertEqual(Student.objects.count(), 0)

        self.client.post(reverse("core:import-commit", args=[batch.pk]),
                         {"reviewed_warnings": "yes"})
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.COMMITTED)

    def test_unexpected_missing_record_fails_the_batch_cleanly(self):
        batch = self.batch_for(workbook([student_row("Asha")]))
        original = imports._apply

        def missing(*args, **kwargs):
            raise Guardian.DoesNotExist

        imports._apply = missing
        try:
            with self.assertRaises(imports.CommitError):
                imports.commit_batch(batch, user=self.admin)
        finally:
            imports._apply = original
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.FAILED)

    def test_discard_and_expiry(self):
        batch = self.batch_for(workbook([student_row("Asha")]))
        self.client.post(reverse("core:import-discard", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.DISCARDED)
        self.assertEqual(batch.payload, {})

        stale = self.batch_for(workbook([student_row("Arjun")]))
        later = timezone.now() + imports.PENDING_LIFETIME + datetime.timedelta(minutes=1)
        self.assertEqual(imports.expire_stale_batches(now=later), 1)
        stale.refresh_from_db()
        self.assertEqual(stale.status, ImportBatch.Status.EXPIRED)
        with self.assertRaises(imports.CommitError):
            imports.commit_batch(stale, user=self.admin)
        self.assertEqual(Student.objects.count(), 0)


class HistoryGuardTests(ImportFixture):
    def test_batches_and_row_results_are_history(self):
        rejected = self.batch_for(workbook([student_row("Asha", cls="Nope")]))
        result = rejected.row_results.first()
        with self.assertRaises(DatabaseError), transaction.atomic():
            ImportBatch.objects.filter(pk=rejected.pk).delete()
        with self.assertRaises(DatabaseError), transaction.atomic():
            ImportBatch.objects.filter(pk=rejected.pk).update(status="COMMITTED")
        with self.assertRaises(DatabaseError), transaction.atomic():
            ImportRowResult.objects.filter(pk=result.pk).update(message="changed")
        with self.assertRaises(DatabaseError), transaction.atomic():
            ImportRowResult.objects.filter(pk=result.pk).delete()

        pending = self.batch_for(workbook([student_row("Asha")]))
        with self.assertRaises(DatabaseError), transaction.atomic():
            ImportBatch.objects.filter(pk=pending.pk).update(file_name="other.xlsx")
        imports.commit_batch(pending, user=self.admin)
        with self.assertRaises(DatabaseError), transaction.atomic():
            ImportBatch.objects.filter(pk=pending.pk).update(status="DISCARDED")


class ImportAuthorizationTests(ImportFixture):
    def test_only_administrators_reach_import_pages(self):
        batch = self.batch_for(workbook([student_row("Asha")]))
        urls = [
            reverse("core:import-list"),
            reverse("core:import-upload"),
            reverse("core:import-template"),
            reverse("core:import-detail", args=[batch.pk]),
            reverse("core:import-issues-csv", args=[batch.pk]),
        ]
        User.objects.create_user(username="teacher", password="pass-12345")
        other = self.client_class()
        other.login(username="teacher", password="pass-12345")
        anonymous = self.client_class()
        for url in urls:
            self.assertEqual(other.get(url).status_code, 403, url)
            self.assertEqual(anonymous.get(url).status_code, 302, url)
            self.assertEqual(self.client.get(url).status_code, 200, url)
        for name in ("core:import-commit", "core:import-discard"):
            url = reverse(name, args=[batch.pk])
            self.assertEqual(other.post(url).status_code, 403)
            self.assertEqual(self.client.get(url).status_code, 405)
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.VALIDATED)

    def test_upload_flow_and_downloads(self):
        response = self.client.get(reverse("core:import-template"))
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Students", book.sheetnames)
        self.assertIn("Classes", book.sheetnames)

        response = self.upload(workbook([student_row("Asha", cls="Nope")]))
        batch = ImportBatch.objects.get()
        self.assertRedirects(response, reverse("core:import-detail", args=[batch.pk]))
        page = self.client.get(reverse("core:import-detail", args=[batch.pk]))
        self.assertContains(page, "does not exist")
        self.assertNotContains(page, "Confirm import")
        csv_text = self.client.get(reverse("core:import-issues-csv", args=[batch.pk])).content
        self.assertIn(b"does not exist", csv_text)

        self.upload(workbook([student_row("Asha")]))
        pending = ImportBatch.objects.filter(status="VALIDATED").get()
        page = self.client.get(reverse("core:import-detail", args=[pending.pk]))
        self.assertContains(page, "Confirm import")
        self.assertContains(page, "Class 5 - A")
